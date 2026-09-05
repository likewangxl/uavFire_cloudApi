# -*- coding: utf-8 -*-
"""生成《智能集群大载重无人机灭火系统》消防领导推广版 PPT(场景功能全景)。

素材: assets/leader_20260711/ (截图会自动裁掉浏览器边框等杂物到 _build/)
输出: docs/智能集群大载重无人机灭火系统-消防领导推广版-场景功能全景-20260711.pptx
"""
import os
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets", "leader_20260711")
BUILD = os.path.join(ASSETS, "_build")
OUT = os.path.join(HERE, "..", "智能集群大载重无人机灭火系统-消防领导推广版-场景功能全景-20260711.pptx")
os.makedirs(BUILD, exist_ok=True)

# ---------- 配色 ----------
NAVY_TOP = (7, 13, 28)
NAVY_BOT = (15, 27, 52)
CARD = RGBColor(0x13, 0x1F, 0x38)
CARD_LIGHT = RGBColor(0x1A, 0x2A, 0x4A)
BORDER = RGBColor(0x2C, 0x40, 0x66)
ORANGE = RGBColor(0xFF, 0x6B, 0x35)
AMBER = RGBColor(0xFF, 0xB0, 0x3A)
CYAN = RGBColor(0x3A, 0xC8, 0xE0)
GREEN = RGBColor(0x3A, 0xDF, 0xA5)
VIOLET = RGBColor(0x9B, 0x8C, 0xFF)
WHITE = RGBColor(0xF5, 0xF8, 0xFF)
GREY = RGBColor(0xA9, 0xB7, 0xD1)
DIM = RGBColor(0x72, 0x82, 0xA0)
FONT = "PingFang SC"

W, H = Inches(13.333), Inches(7.5)
IN = Inches


# ---------- 素材预处理:裁掉浏览器边框 / 系统菜单 / 调试按钮 ----------
def prep(src, dst, top=0, bottom=0, left=0, right=0):
    p = os.path.join(BUILD, dst)
    im = Image.open(os.path.join(ASSETS, src))
    w, h = im.size
    im.crop((left, top, w - right, h - bottom)).convert("RGB").save(p, quality=92)
    return p


IMG = {
    "cockpit": prep("cockpit.jpg", "cockpit.jpg", bottom=48),           # 去 vConsole
    "visible": prep("visible.jpg", "visible.jpg"),
    "thermal": prep("thermal.jpg", "thermal.jpg"),
    "thermal_night": prep("thermal_night.png", "thermal_night.jpg"),
    "incident_map": prep("incident_map.jpg", "incident_map.jpg", bottom=48),
    "event_list": prep("event_list.jpg", "event_list.jpg", top=58, bottom=40),
    "task_create": prep("task_create.jpg", "task_create.jpg", top=58, bottom=40),
    "plan_nfz": prep("plan_nfz.jpg", "plan_nfz.jpg", top=58, bottom=40),
    "save_block": prep("save_block.jpg", "save_block.jpg", top=58, bottom=40),
    "dual_cockpit": prep("dual_cockpit.jpg", "dual_cockpit.jpg", top=92, bottom=72),
    "fc100_payload": prep("fc100_payload.jpg", "fc100_payload.jpg"),
    "suppression": prep("suppression.jpg", "suppression.jpg"),
}


# ---------- PIL 背景 ----------
def _glow(img, center, radius, color, alpha):
    ov = Image.new("RGBA", img.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(ov)
    d.ellipse([center[0] - radius, center[1] - radius,
               center[0] + radius, center[1] + radius], fill=color + (alpha,))
    img.alpha_composite(ov.filter(ImageFilter.GaussianBlur(radius // 2)))


def _lines(img, draw_fn):
    ov = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw_fn(ImageDraw.Draw(ov))
    img.alpha_composite(ov)


def make_bg(path, kind):
    w, h = 1920, 1080
    img = Image.new("RGBA", (w, h))
    d = ImageDraw.Draw(img)
    for y in range(h):
        t = y / h
        c = tuple(int(NAVY_TOP[i] + (NAVY_BOT[i] - NAVY_TOP[i]) * t) for i in range(3))
        d.line([(0, y), (w, y)], fill=c + (255,))
    if kind == "cover":
        _glow(img, (1660, 900), 560, (255, 95, 40), 78)
        _glow(img, (150, 90), 430, (40, 170, 220), 50)
        def deco(dd):
            for rad in (330, 420, 510, 600):
                dd.ellipse([1660 - rad, 900 - rad, 1660 + rad, 900 + rad],
                           outline=(255, 140, 70, 40), width=2)
            for x in range(0, w, 96):
                dd.line([(x, 0), (x, h)], fill=(255, 255, 255, 7))
        _lines(img, deco)
    elif kind == "section":
        _glow(img, (1780, 1020), 460, (255, 95, 40), 46)
        _glow(img, (120, 60), 360, (40, 170, 220), 34)
        def deco(dd):
            for rad in (260, 340):
                dd.ellipse([1780 - rad, 1020 - rad, 1780 + rad, 1020 + rad],
                           outline=(255, 140, 70, 36), width=2)
        _lines(img, deco)
    else:
        _glow(img, (1830, 1050), 340, (255, 95, 40), 26)
        _glow(img, (80, 40), 280, (40, 170, 220), 20)
    img.convert("RGB").save(path, quality=92)
    return path


BG_COVER = make_bg(os.path.join(BUILD, "bg_cover.jpg"), "cover")
BG_PAGE = make_bg(os.path.join(BUILD, "bg_page.jpg"), "page")
BG_SECTION = make_bg(os.path.join(BUILD, "bg_section.jpg"), "section")


# ---------- pptx 工具 ----------
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
PAGENO = [0]


def set_font(run, size, color=WHITE, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", FONT)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=1.12):
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(size * 0.30)
        run = p.add_run()
        run.text = text.replace(",", "，").replace(";", "；").replace(":", "：")
        set_font(run, size, color, bold)
    return tb


def panel(slide, x, y, w, h, fill=CARD, line=BORDER, rounded=True, radius=0.055):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        IN(x), IN(y), IN(w), IN(h))
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.0)
    return shp


def accent_bar(slide, x, y, w=0.55, h=0.05, color=ORANGE):
    panel(slide, x, y, w, h, fill=color, line=None, rounded=False)


def header(slide, chapter, chip=None):
    PAGENO[0] += 1
    panel(slide, 0.65, 0.475, 0.09, 0.09, fill=ORANGE, line=None, rounded=False)
    add_text(slide, 0.84, 0.40, 7.5, 0.26, [(chapter, 10.5, DIM, False)])
    add_text(slide, 11.6, 0.40, 1.05, 0.26,
             [("%02d" % PAGENO[0], 10.5, DIM, False)], align=PP_ALIGN.RIGHT)
    if chip:
        cw = 0.42 + len(chip) * 0.165
        panel(slide, 12.68 - cw, 0.78, cw, 0.34, fill=CARD_LIGHT, line=BORDER)
        add_text(slide, 12.68 - cw, 0.845, cw, 0.24,
                 [(chip, 10.5, AMBER, True)], align=PP_ALIGN.CENTER)


def title_lead(slide, title, lead=None, tw=11.0, lw=11.9):
    add_text(slide, 0.65, 0.86, tw, 0.55, [(title, 24, WHITE, True)])
    accent_bar(slide, 0.66, 1.46)
    if lead:
        add_text(slide, 0.65, 1.66, lw, 0.85, [(lead, 12.5, GREY, False)], spacing=1.3)


PICS = []  # (page_idx, x, y, w, h) 用于遮挡自检


def pic_card(slide, key, x, y, max_w, max_h, caption=None, align="left"):
    """图片完整等比放入 (max_w, max_h),外框贴合图片,说明文字放在图片下方(不遮挡)。"""
    path = IMG[key]
    iw, ih = Image.open(path).size
    pad = 0.07
    aw, ah = max_w - 2 * pad, max_h - 2 * pad
    scale = min(aw / iw, ah / ih)
    fw, fh = iw * scale, ih * scale
    if align == "center":
        px = x + (max_w - fw - 2 * pad) / 2
    elif align == "right":
        px = x + (max_w - fw - 2 * pad)
    else:
        px = x
    panel(slide, px, y, fw + 2 * pad, fh + 2 * pad, fill=CARD, line=BORDER, radius=0.03)
    slide.shapes.add_picture(path, IN(px + pad), IN(y + pad), IN(fw), IN(fh))
    PICS.append((PAGENO[0], px, y, fw + 2 * pad, fh + 2 * pad))
    bottom = y + fh + 2 * pad
    if caption:
        panel(slide, px + 0.02, bottom + 0.13, 0.05, 0.15, fill=ORANGE, line=None, rounded=False)
        add_text(slide, px + 0.16, bottom + 0.10, fw - 0.2, 0.24,
                 [(caption, 10, DIM, False)])
        bottom += 0.38
    return bottom


def bullets(slide, x, y, w, items, gap=0.30, tsize=13, dsize=11):
    cy = y
    for t, d in items:
        panel(slide, x, cy + 0.075, 0.10, 0.10, fill=ORANGE, line=None, rounded=False)
        add_text(slide, x + 0.24, cy, w - 0.24, 0.3, [(t, tsize, WHITE, True)])
        if d:
            add_text(slide, x + 0.24, cy + 0.32, w - 0.24, 0.55,
                     [(d, dsize, GREY, False)], spacing=1.22)
            cy += 0.32 + 0.30 * (1 + (len(d) > 22)) + gap
        else:
            cy += 0.32 + gap
    return cy


def num_badge(slide, x, y, num, color=ORANGE, size=0.36, fsize=12):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, IN(x), IN(y), IN(size), IN(size))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD_LIGHT
    shp.line.color.rgb = color
    shp.line.width = Pt(1.2)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = num
    set_font(r, fsize, color, True)


def bottom_band(slide, text, y=6.62, color=AMBER, w=12.03):
    panel(slide, 0.65, y, w, 0.52, fill=CARD, line=BORDER)
    add_text(slide, 0.9, y + 0.13, w - 0.5, 0.3, [(text, 12.5, color, True)],
             align=PP_ALIGN.CENTER)


def bg(slide, path=BG_PAGE):
    slide.shapes.add_picture(path, 0, 0, W, H)


def new_slide(bgpath=BG_PAGE):
    s = prs.slides.add_slide(BLANK)
    bg(s, bgpath)
    return s


# =========================================================
# P1 封面
# =========================================================
s = new_slide(BG_COVER)
PAGENO[0] += 1
add_text(s, 0.85, 0.92, 9.0, 0.3, [("UAVFIRE · 面向消防救援的空地一体化作业平台", 12.5, CYAN, True)])
accent_bar(s, 0.86, 1.34, w=0.62)
add_text(s, 0.85, 1.86, 11.6, 2.4, [
    ("智能集群大载重", 40, WHITE, True),
    ("无人机灭火系统", 40, WHITE, True),
], spacing=1.08)
add_text(s, 0.85, 4.28, 7.9, 1.0, [
    ("打通侦察发现、智能研判、重载处置的完整链路——用一套系统,把“看火的机”与“灭火的机”组织成一支协同作战的空中力量。",
     13.5, GREY, False)], spacing=1.35)
chips = ["森林消防", "草原消防", "大型仓储消防", "应急救援", "物资投放"]
cx = 0.85
for c in chips:
    cw = 0.5 + len(c) * 0.21
    panel(s, cx, 5.72, cw, 0.44, fill=CARD_LIGHT, line=BORDER)
    add_text(s, cx, 5.815, cw, 0.26, [(c, 11.5, WHITE, False)], align=PP_ALIGN.CENTER)
    cx += cw + 0.22
add_text(s, 0.85, 6.72, 6.0, 0.28, [("2026 年 7 月 · 汇报材料", 10.5, DIM, False)])

# =========================================================
# P2 目录
# =========================================================
s = new_slide(BG_SECTION)
header(s, "汇报内容")
add_text(s, 0.65, 0.98, 11.0, 0.55, [("从系统到场景,从能力到实战", 25, WHITE, True)])
accent_bar(s, 0.66, 1.60)
toc = [
    ("01", "系统概述", "平台定位,与“侦察感知—指挥中枢—重载处置”三层体系"),
    ("02", "应用场景", "森林、草原、大型仓储消防,以及应急救援、物资投放"),
    ("03", "核心特点", "跨行业融合、多模态识别、适航管理等六项核心能力"),
    ("04", "实战与总结", "一次完整接力的全流程,以及系统带来的三层价值"),
]
ty = 2.15
for num, t, d in toc:
    panel(s, 0.65, ty, 12.03, 1.08, fill=CARD, line=BORDER)
    add_text(s, 1.05, ty + 0.26, 1.1, 0.55, [(num, 26, ORANGE, True)])
    panel(s, 2.25, ty + 0.24, 0.016, 0.60, fill=BORDER, line=None, rounded=False)
    add_text(s, 2.62, ty + 0.17, 3.2, 0.4, [(t, 17, WHITE, True)])
    add_text(s, 2.62, ty + 0.60, 9.7, 0.32, [(d, 11.5, GREY, False)])
    ty += 1.26

# =========================================================
# P3 系统概述
# =========================================================
s = new_slide()
header(s, "01 · 系统概述")
title_lead(
    s, "一个平台,统管“看火的机”与“灭火的机”",
    "系统面向消防救援实战,把两类原本分属不同行业体系的无人机——行业侦察机与大载重运载机——纳入同一套指挥平台:"
    "侦察机负责看得早、辨得准,运载机负责到得快、投得准,指挥员在一块屏幕上完成从发现火情到组织处置的全部动作。")
roles = [
    ("侦察感知层", "M4T 行业侦察机", "可见光+热成像双光云台,昼夜自动巡查;AI 实时识别烟、火与异常热源。", CYAN),
    ("指挥中枢层", "统一指挥平台", "态势一屏统览;事件、航线、任务与权限统一调度,全程数据留痕。", AMBER),
    ("重载处置层", "FC100 大载重运载机", "数十公斤级运力,灭火弹与救援物资精准投送,到点自动脱钩。", ORANGE),
]
rx = 0.65
for tag, name, desc, color in roles:
    panel(s, rx, 2.88, 3.85, 2.45, fill=CARD, line=BORDER)
    panel(s, rx, 2.88, 0.07, 2.45, fill=color, line=None, rounded=False)
    add_text(s, rx + 0.32, 3.12, 3.3, 0.28, [(tag, 11, color, True)])
    add_text(s, rx + 0.32, 3.48, 3.3, 0.36, [(name, 16, WHITE, True)])
    add_text(s, rx + 0.32, 3.98, 3.25, 1.2, [(desc, 11.5, GREY, False)], spacing=1.3)
    rx += 4.04
bottom_band(s, "侦察机发现火情  →  平台生成任务  →  运载机接力处置:三层能力在同一条链路上协同", y=5.72)
add_text(s, 0.65, 6.55, 12.0, 0.6,
         [("对用户而言,这意味着:不再为不同机型部署不同系统、培训不同终端——一个平台,就是一支完整的空中消防分队。",
           11.5, DIM, False)], align=PP_ALIGN.CENTER)

# =========================================================
# P4 全链路闭环
# =========================================================
s = new_slide()
header(s, "01 · 系统概述")
title_lead(
    s, "一次火情,从发现到复盘的完整闭环",
    "灭火不是单个动作,而是一串环环相扣的决策。系统把六个环节固化为标准作业链——每一步的输入都来自上一步的输出,不断链、不丢信息。")
steps = [
    ("巡查发现", "侦察机按预案航线自动巡查,画面实时回传"),
    ("智能研判", "AI 对可见光与热成像画面融合识别、风险分级"),
    ("事件确认", "指挥员复核证据,确认火情、审批处置"),
    ("航线生成", "系统按火点坐标自动生成投放航线"),
    ("重载处置", "运载机自主抵达,到点投放自动脱钩"),
    ("复查归档", "侦察机复查现场,事件全程归档可溯"),
]
sx = 0.65
for i, (t, d) in enumerate(steps):
    shp = s.shapes.add_shape(
        MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON,
        IN(sx), IN(2.92), IN(2.1), IN(0.68))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD_LIGHT if i % 2 == 0 else CARD
    shp.line.color.rgb = BORDER
    shp.line.width = Pt(1.0)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = t
    set_font(r, 13, AMBER if i in (1, 4) else WHITE, True)
    add_text(s, sx + 0.16, 3.82, 1.68, 1.5, [(d, 10.5, GREY, False)], spacing=1.25)
    num_badge(s, sx + 0.86, 5.35, "%02d" % (i + 1), color=DIM, size=0.32, fsize=10.5)
    sx += 1.97
bottom_band(s, "链条上每一个动作都有系统记录:谁确认、何时下发、执行结果如何,事后都能完整回放", y=6.25, color=GREY)
add_text(s, 0.65, 6.98, 12.0, 0.3,
         [("其中“智能研判”与“重载处置”两环,是本系统区别于普通巡检无人机的关键。", 11, DIM, False)],
         align=PP_ALIGN.CENTER)

# =========================================================
# P5 应用场景总览
# =========================================================
s = new_slide(BG_SECTION)
header(s, "02 · 应用场景")
title_lead(
    s, "一套系统,覆盖“防、灭、救、送”四类任务",
    "平台的底层能力是“空中发现 + 空中处置”。凡是需要先看清、再抵近、能投送、可复查的现场,都是它的用武之地——以下五类场景均已形成明确打法。")
scenes_top = [
    ("01", "森林消防", "防火期无人值守巡护,火情早发现、初期快处置,不给山火成势的机会。", ORANGE),
    ("02", "草原消防", "大范围高频巡查,弥补人力与车辆覆盖半径不足,火烧迹地快速复查。", AMBER),
    ("03", "大型仓储消防", "厂房、库区与堆场的热成像巡测,阴燃早预警,重点目标定期建档。", CYAN),
]
scenes_bot = [
    ("04", "应急救援", "山地、洪涝、震后等复杂现场先期侦察,热成像搜寻热源与人员线索。", GREEN),
    ("05", "物资投放", "灭火弹、药剂、食品与救援装备的重载精准投送,到点自动脱钩。", VIOLET),
]
rx = 0.65
for num, t, d, color in scenes_top:
    panel(s, rx, 2.72, 3.85, 1.88, fill=CARD, line=BORDER)
    add_text(s, rx + 0.28, 2.92, 1.0, 0.4, [(num, 19, color, True)])
    add_text(s, rx + 0.95, 2.97, 2.7, 0.34, [(t, 15.5, WHITE, True)])
    add_text(s, rx + 0.28, 3.48, 3.32, 1.0, [(d, 11, GREY, False)], spacing=1.28)
    rx += 4.04
rx = 0.65
for num, t, d, color in scenes_bot:
    panel(s, rx, 4.82, 5.87, 1.52, fill=CARD, line=BORDER)
    add_text(s, rx + 0.28, 5.02, 1.0, 0.4, [(num, 19, color, True)])
    add_text(s, rx + 0.95, 5.07, 3.4, 0.34, [(t, 15.5, WHITE, True)])
    add_text(s, rx + 0.28, 5.58, 5.35, 0.62, [(d, 11, GREY, False)], spacing=1.28)
    rx += 6.06
bottom_band(s, "场景在变,链路不变:发现 → 研判 → 任务 → 处置 → 复查", y=6.62)

# =========================================================
# P6 场景 · 森林草原
# =========================================================
s = new_slide()
header(s, "02 · 应用场景", chip="场景 · 森林 / 草原")
add_text(s, 0.65, 0.86, 10.2, 0.55, [("把巡护搬上天空,把火情灭在初期", 24, WHITE, True)])
accent_bar(s, 0.66, 1.46)
add_text(s, 0.65, 1.72, 5.5, 1.9, [
    ("防火期内,侦察机按预案航线昼夜自动巡查林区与草场,AI 持续识别烟、火与异常热源;一旦确认火情,运载机即携灭火弹抵近投放——在火势蔓延之前完成“早发现、早处置”。",
     12.5, GREY, False)], spacing=1.35)
bullets(s, 0.65, 3.78, 5.5, [
    ("无人值守巡护", "预案航线周期起飞,大范围覆盖不靠人盯"),
    ("昼夜双光值守", "白天识烟火,夜间锁热源,防火期全时段在线"),
    ("初期快速压制", "发现即定位、定位即投放,不给火势留时间"),
], gap=0.26)
pic_card(s, "suppression", 6.42, 2.35, 6.26, 3.9,
         caption="灭火弹投放后,火点即刻压制(外场实拍)", align="right")
bottom_band(s, "对森林草原而言,“早十分钟发现、早十分钟处置”,就是一场大火与一次记录的差别", y=6.62)

# =========================================================
# P7 场景 · 大型仓储
# =========================================================
s = new_slide()
header(s, "02 · 应用场景", chip="场景 · 大型仓储")
add_text(s, 0.65, 0.86, 10.2, 0.55, [("盯住每一处阴燃,在起火之前预警", 24, WHITE, True)])
accent_bar(s, 0.66, 1.46)
pic_card(s, "thermal_night", 0.65, 2.35, 6.26, 3.9,
         caption="夜间热成像自动锁定阴燃热源,并标注 153.0℃ 温度读数")
add_text(s, 7.25, 1.72, 5.43, 1.9, [
    ("大型仓储、物流园区与重点库区的火灾多起于阴燃,肉眼和普通监控难以察觉。热成像巡测能直接读出屋面与堆场的温度异常——一处发热、立即告警建档,为处置争取最宝贵的窗口期。",
     12.5, GREY, False)], spacing=1.35)
bullets(s, 7.25, 3.78, 5.43, [
    ("夜间自动巡测", "无需照明,热源在黑暗中一览无余"),
    ("温度读数留证", "异常点位自动标注温度并截图归档"),
    ("告警直达值守", "达到阈值即时推送,从发现到派单一步到位"),
], gap=0.26)
bottom_band(s, "同样的打法,可直接复用于化工园区、煤场、垃圾场、电力设施等重点防火目标", y=6.62)

# =========================================================
# P8 场景 · 救援与物资投放
# =========================================================
s = new_slide()
header(s, "02 · 应用场景", chip="场景 · 救援 / 投放")
title_lead(
    s, "到不了的地方,让无人机先到",
    "山地、洪涝、震后等现场,地面力量往往难以第一时间抵近。大载重运载机具备数十公斤级运力,可先行侦察、再行投送——把灭火弹、药剂、食品与救援装备,准确送到需要的位置。")
pic_card(s, "fc100_payload", 0.65, 2.72, 6.9, 3.68,
         caption="FC100 运载机吊挂灭火载荷,飞赴目标区(外场实拍)")
cards = [
    ("先期侦察", "热成像搜寻热源与人员线索,先于队伍回传现场态势"),
    ("重载投送", "灭火弹、药剂、救生装备一次挂载、精准送达"),
    ("到点自投", "航点到达自动脱钩投放,全程人员零涉险"),
]
cy = 2.72
for i, (t, d) in enumerate(cards):
    panel(s, 7.85, cy, 4.83, 1.06, fill=CARD, line=BORDER)
    num_badge(s, 8.08, cy + 0.34, "%d" % (i + 1), color=ORANGE, size=0.36)
    add_text(s, 8.62, cy + 0.15, 3.9, 0.32, [(t, 13.5, WHITE, True)])
    add_text(s, 8.62, cy + 0.50, 3.9, 0.5, [(d, 10.5, GREY, False)], spacing=1.2)
    cy += 1.26
bottom_band(s, "“救”与“送”用的是与灭火完全相同的系统能力——多一类场景,不多一套系统", y=6.82)

# =========================================================
# P9 核心特点总览
# =========================================================
s = new_slide(BG_SECTION)
header(s, "03 · 核心特点")
title_lead(
    s, "六项核心能力,支撑一次完整的空中处置",
    "这些能力不是功能罗列,而是围绕“发现—研判—处置”主线逐环设计:缺任何一环,链路就会断在人工衔接上。")
feats = [
    ("①", "跨行业融合 · 任务接续", "侦察机与运载机同平台纳管,火情事件直接转为处置任务"),
    ("②", "多模态火情识别", "可见光+热成像+AI 融合研判,风险分级、人工把关"),
    ("③", "适航飞行 · 限飞区管理", "禁飞区规则前置校验,违规航线在起飞前被拦截"),
    ("④", "全自主任务链", "规划即下发,起飞、飞行、投放、返航全程自主执行"),
    ("⑤", "无人值守 · 一屏指挥", "系统替人盯屏、异常唤醒,态势与任务同屏调度"),
    ("⑥", "精准定位 · 全程留痕", "火点解算为带误差半径的真实坐标,事件档案可复盘"),
]
for i, (num, t, d) in enumerate(feats):
    fx = 0.65 + (i % 3) * 4.04
    fy = 2.78 + (i // 3) * 2.06
    panel(s, fx, fy, 3.85, 1.84, fill=CARD, line=BORDER)
    add_text(s, fx + 0.26, fy + 0.18, 0.6, 0.44, [(num, 20, ORANGE, True)])
    add_text(s, fx + 0.82, fy + 0.24, 2.95, 0.36, [(t, 13.5, WHITE, True)])
    add_text(s, fx + 0.26, fy + 0.78, 3.35, 0.9, [(d, 10.5, GREY, False)], spacing=1.28)
add_text(s, 0.65, 6.90, 12.03, 0.3, [("——以下六页,逐项展开——", 11, DIM, False)],
         align=PP_ALIGN.CENTER)

# =========================================================
# P10 特点① 跨行业融合·任务接续
# =========================================================
s = new_slide()
header(s, "03 · 核心特点", chip="特点 ① 跨行业融合 · 任务接续")
title_lead(
    s, "侦察机发现的火情,直接变成运载机的任务",
    "行业侦察机与大载重运载机分属不同行业体系,原本各用各的地面站与终端。平台把两类机型统一接入、统一调度,"
    "火情事件在系统内直接流转为投放任务——坐标不需要人工转述,任务不需要跨系统倒手。")
flow = ["侦察机发现", "平台建事件", "生成投放航线", "运载机接力执行"]
fx = 0.65
for i, t in enumerate(flow):
    fw = 2.55 if i < 3 else 2.75
    panel(s, fx, 2.78, fw, 0.5, fill=CARD_LIGHT, line=BORDER)
    add_text(s, fx, 2.885, fw, 0.3, [(t, 12.5, AMBER if i in (1, 3) else WHITE, True)],
             align=PP_ALIGN.CENTER)
    if i < 3:
        add_text(s, fx + fw + 0.02, 2.855, 0.5, 0.32, [("→", 14, ORANGE, True)],
                 align=PP_ALIGN.CENTER)
    fx += fw + 0.52
pic_card(s, "dual_cockpit", 0.65, 3.62, 6.55, 3.15,
         caption="巡检机与运载机在同一驾驶舱内协同调度(外场实飞画面)")
bullets(s, 7.55, 3.86, 5.13, [
    ("异构机队统一接入", "一套界面同时管“看火的机”与“灭火的机”"),
    ("事件驱动任务生成", "从火情告警到运载机领任务,分钟级完成"),
    ("接力全程状态回传", "抵近、待命、投放、返航,指挥员实时掌握"),
], gap=0.28)

# =========================================================
# P11 特点② 多模态识别
# =========================================================
s = new_slide()
header(s, "03 · 核心特点", chip="特点 ② 多模态火情识别")
title_lead(
    s, "同一现场,两种“眼睛”,一次研判",
    "可见光相机捕捉烟雾与明火,红外热成像穿透夜色与烟障锁定温度异常;AI 对两路画面同时识别、融合打分,"
    "原始证据全部保留——复杂环境下不漏报、少误报,复核与追溯都有依据。")
pic_card(s, "visible", 0.65, 2.72, 5.87, 3.5, caption="可见光:识别烟雾与疑似火点")
pic_card(s, "thermal", 6.81, 2.72, 5.87, 3.5, caption="热成像:同一现场锁定 153℃ 异常热源")
bottom_band(s, "可见光识别 ＋ 热成像测温 ＋ AI 融合分级 ＋ 人工最终确认 —— 四道关口,一次研判", y=6.62)

# =========================================================
# P12 特点③ 适航飞行·限飞区管理
# =========================================================
s = new_slide()
header(s, "03 · 核心特点", chip="特点 ③ 适航飞行 · 限飞区管理")
title_lead(
    s, "合规不靠自觉,系统在起飞前把关",
    "平台内置禁飞区与适飞空域规则:航线规划阶段,冲突航点与穿越航段实时标红;保存与下发前强制提示,"
    "是否调整、是否放行,由有权限的指挥员决定——“能不能飞”的判断,被前置到起飞之前。")
pic_card(s, "plan_nfz", 0.65, 2.66, 7.35, 3.75,
         caption="规划页实时标出禁飞区冲突航点与穿越航段,并同步给出地形高度剖面")
pic_card(s, "save_block", 8.15, 2.66, 4.53, 2.6,
         caption="保存时强制拦截:返回修改,或授权放行")
bullets(s, 8.15, 5.62, 4.53, [
    ("规划期实时校验", None),
    ("保存 / 下发双重闸口", None),
    ("权限化放行,全程留痕", None),
], gap=0.115, tsize=12.5)

# =========================================================
# P13 特点④ 全自主任务链
# =========================================================
s = new_slide()
header(s, "03 · 核心特点", chip="特点 ④ 全自主任务链")
title_lead(
    s, "规划即下发,从起飞到返航无需逐步操控",
    "指挥员在地图上设计好航线、一键下发,无人机自动起飞、逐航点飞行、按编排执行拍照、测温、投放等动作,"
    "完成后自动返航,运载机到点自动脱钩。飞手从“全程操控”解放为“关键确认”。")
chain = [
    ("一键下发", "航线与动作编排一次下发到机"),
    ("自动起飞", "无需人工推杆,自主离场入航"),
    ("航线飞行", "逐航点按设定高度速度飞行"),
    ("动作执行", "拍照 · 测温 · 投放脱钩自动完成"),
    ("自动返航", "任务完成自主返航降落"),
]
cy = 2.72
for i, (t, d) in enumerate(chain):
    panel(s, 0.65, cy, 4.55, 0.62, fill=CARD, line=BORDER)
    num_badge(s, 0.84, cy + 0.13, "%d" % (i + 1), color=CYAN, size=0.36)
    add_text(s, 1.38, cy + 0.06, 1.55, 0.3, [(t, 12.5, WHITE, True)])
    add_text(s, 1.38, cy + 0.345, 3.7, 0.26, [(d, 9.5, DIM, False)])
    if i < 4:
        panel(s, 1.0, cy + 0.62, 0.016, 0.16, fill=BORDER, line=None, rounded=False)
    cy += 0.78
pic_card(s, "task_create", 5.62, 2.72, 7.06, 3.75,
         caption="监测任务与投放任务在同一入口创建、同一航线库管理", align="right")

# =========================================================
# P14 特点⑤ 无人值守·一屏指挥
# =========================================================
s = new_slide()
header(s, "03 · 核心特点", chip="特点 ⑤ 无人值守 · 一屏指挥")
title_lead(
    s, "系统替人盯屏,一块屏幕指挥全局",
    "值守模式下,平台按预案周期派飞、自动识别、异常唤醒——值守人员从“盯屏幕”变为“等告警”。"
    "指挥驾驶舱把火情、机队、直播画面与任务进度汇聚一屏,判断与调度不再切换系统。")
pic_card(s, "cockpit", 0.65, 2.62, 9.05, 4.45,
         caption="指挥驾驶舱:火情、机队、直播与任务,一屏统览")
chips5 = [("预案周期巡航", "按计划自动起飞巡查"),
          ("异常主动告警", "识别到风险即时唤醒值守"),
          ("态势任务同屏", "判断与调度零切换")]
cy = 2.74
for t, d in chips5:
    panel(s, 10.05, cy, 2.63, 1.06, fill=CARD, line=BORDER)
    add_text(s, 10.25, cy + 0.16, 2.3, 0.3, [(t, 12.5, AMBER, True)])
    add_text(s, 10.25, cy + 0.52, 2.3, 0.5, [(d, 10, GREY, False)], spacing=1.18)
    cy += 1.26

# =========================================================
# P15 特点⑥ 精准定位·全程留痕
# =========================================================
s = new_slide()
header(s, "03 · 核心特点", chip="特点 ⑥ 精准定位 · 全程留痕")
title_lead(
    s, "报出的是坐标,留下的是证据",
    "系统融合定位、姿态、云台角度与地形高程,把画面中的火点解算为带误差半径的真实坐标,直接用于生成投放航线;"
    "识别快照、人工确认、审批与处置动作全部入档——每一次响应都可回放、可复盘。")
pic_card(s, "incident_map", 0.65, 2.72, 6.75, 3.35,
         caption="火点坐标与风险半径直接标注在处置地图上,坐标质量一目了然")
pic_card(s, "event_list", 7.6, 3.17, 5.08, 3.35,
         caption="识别快照、置信度、人工确认与审批,形成完整事件档案")
bottom_band(s, "坐标可用于行动,档案可用于复盘 —— 数据既是“弹药”,也是“证据”", y=6.62)

# =========================================================
# P16 实战闭环
# =========================================================
s = new_slide(BG_SECTION)
header(s, "04 · 实战与总结")
title_lead(
    s, "一次完整的接力:从发现火情到确认扑灭",
    "以下六步已在系统内完整贯通,并经外场实飞验证——全程在同一平台内完成,不需要切换任何第三方系统。")
tl = [
    ("01", "自动巡查", "侦察机按预案航线起飞巡护"),
    ("02", "AI 告警", "识别烟火热源,自动生成火情事件"),
    ("03", "人工确认", "复核画面与坐标,审批处置方案"),
    ("04", "航线生成", "按火点坐标一键生成投放航线"),
    ("05", "重载投放", "运载机自主抵达,到点脱钩灭火"),
    ("06", "复查归档", "回场复查效果,事件闭环存档"),
]
panel(s, 0.9, 4.24, 11.53, 0.022, fill=BORDER, line=None, rounded=False)
for i, (num, t, d) in enumerate(tl):
    cx = 1.4 + i * 2.1
    num_badge(s, cx - 0.21, 4.04, num, color=ORANGE if i in (1, 4) else CYAN, size=0.42, fsize=12)
    if i % 2 == 0:
        add_text(s, cx - 1.0, 3.06, 2.0, 0.32, [(t, 13.5, WHITE, True)], align=PP_ALIGN.CENTER)
        add_text(s, cx - 1.0, 3.40, 2.0, 0.55, [(d, 10, GREY, False)],
                 align=PP_ALIGN.CENTER, spacing=1.2)
    else:
        add_text(s, cx - 1.0, 4.72, 2.0, 0.32, [(t, 13.5, WHITE, True)], align=PP_ALIGN.CENTER)
        add_text(s, cx - 1.0, 5.06, 2.0, 0.55, [(d, 10, GREY, False)],
                 align=PP_ALIGN.CENTER, spacing=1.2)
bottom_band(s, "从 AI 告警到运载机起飞,分钟级完成——而每一步,都留有记录", y=6.32)
add_text(s, 0.65, 7.02, 12.0, 0.3,
         [("这不是六个功能的演示,而是一套已经跑通的空中处置战法。", 11, DIM, False)],
         align=PP_ALIGN.CENTER)

# =========================================================
# P17 总结
# =========================================================
s = new_slide(BG_SECTION)
header(s, "04 · 实战与总结")
title_lead(
    s, "总结:看得早、到得快、管得住",
    "回到消防救援最朴素的三个诉求,系统给出的答案是:")
vals = [
    ("看得早", "无人值守巡护 + 多模态识别,把发现时点提前到火情初期;夜间与复杂环境同样在线。", CYAN),
    ("到得快", "事件直转任务、双机接力处置,从告警到起飞分钟级;重载运力让处置不止于“看”。", ORANGE),
    ("管得住", "适航校验、权限审批、全程留痕——每一次起飞都合规,每一步动作都可追溯。", AMBER),
]
rx = 0.65
for t, d, color in vals:
    panel(s, rx, 2.62, 3.85, 2.2, fill=CARD, line=BORDER)
    panel(s, rx, 2.62, 3.85, 0.07, fill=color, line=None, rounded=False)
    add_text(s, rx + 0.3, 2.92, 3.2, 0.44, [(t, 19, color, True)])
    add_text(s, rx + 0.3, 3.52, 3.28, 1.2, [(d, 11.5, GREY, False)], spacing=1.32)
    rx += 4.04
panel(s, 0.65, 5.22, 12.03, 0.78, fill=CARD, line=BORDER)
add_text(s, 0.9, 5.40, 11.5, 0.4, [
    ("同一套平台,服务森林、草原、大型仓储消防,以及应急救援、物资投放五类场景", 13.5, WHITE, True)],
    align=PP_ALIGN.CENTER)
add_text(s, 0.65, 6.40, 12.03, 0.4, [
    ("空地一体 · 侦灭一体 —— 让无人机从“会看火”,走到“能灭火”", 15, AMBER, True)],
    align=PP_ALIGN.CENTER)

# =========================================================
# P18 尾页
# =========================================================
s = new_slide(BG_COVER)
PAGENO[0] += 1
add_text(s, 0.85, 2.35, 11.6, 0.35, [("UAVFIRE · 智能集群大载重无人机灭火系统", 13, CYAN, True)],
         align=PP_ALIGN.CENTER)
add_text(s, 0.85, 3.05, 11.6, 0.8, [("让天空成为消防救援的第一响应力量", 32, WHITE, True)],
         align=PP_ALIGN.CENTER)
accent_bar(s, 6.37, 4.05, w=0.6)
add_text(s, 0.85, 4.45, 11.6, 0.4, [("汇报完毕,恳请各位领导指导", 14, GREY, False)],
         align=PP_ALIGN.CENTER)
add_text(s, 0.85, 6.72, 11.6, 0.3, [("2026 年 7 月", 10.5, DIM, False)],
         align=PP_ALIGN.CENTER)

prs.save(OUT)
print("saved:", os.path.abspath(OUT), "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
